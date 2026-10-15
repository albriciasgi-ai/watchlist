"""
Script para crear presentación PowerPoint sobre validación de xStocks
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

# Crear presentación
prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(7.5)

# Definir colores corporativos
COLOR_PRIMARY = RGBColor(31, 78, 120)  # Azul oscuro
COLOR_SUCCESS = RGBColor(76, 175, 80)  # Verde
COLOR_WARNING = RGBColor(255, 152, 0)  # Naranja
COLOR_DANGER = RGBColor(244, 67, 54)   # Rojo
COLOR_TEXT = RGBColor(33, 33, 33)      # Gris oscuro

def add_title_slide(prs, title, subtitle):
    """Slide de título"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank

    # Fondo
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = COLOR_PRIMARY

    # Título
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(9), Inches(1))
    title_frame = title_box.text_frame
    title_frame.text = title
    title_p = title_frame.paragraphs[0]
    title_p.alignment = PP_ALIGN.CENTER
    title_p.font.size = Pt(54)
    title_p.font.bold = True
    title_p.font.color.rgb = RGBColor(255, 255, 255)

    # Subtítulo
    subtitle_box = slide.shapes.add_textbox(Inches(0.5), Inches(4), Inches(9), Inches(0.8))
    subtitle_frame = subtitle_box.text_frame
    subtitle_frame.text = subtitle
    subtitle_p = subtitle_frame.paragraphs[0]
    subtitle_p.alignment = PP_ALIGN.CENTER
    subtitle_p.font.size = Pt(24)
    subtitle_p.font.color.rgb = RGBColor(200, 200, 200)

    return slide

def add_content_slide(prs, title):
    """Slide de contenido con título"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank

    # Título
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.6))
    title_frame = title_box.text_frame
    title_frame.text = title
    title_p = title_frame.paragraphs[0]
    title_p.font.size = Pt(32)
    title_p.font.bold = True
    title_p.font.color.rgb = COLOR_PRIMARY

    # Línea separadora
    line = slide.shapes.add_shape(1, Inches(0.5), Inches(1), Inches(9), Inches(0.01))
    line.fill.solid()
    line.fill.fore_color.rgb = COLOR_PRIMARY
    line.line.fill.background()

    return slide

# ============================================================================
# SLIDE 1: Portada
# ============================================================================
add_title_slide(prs,
    "Validación API xStocks Bybit",
    "Análisis de Compatibilidad y Resultados | Octubre 2026"
)

# ============================================================================
# SLIDE 2: Resumen Ejecutivo
# ============================================================================
slide = add_content_slide(prs, "Resumen Ejecutivo")

content_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(9), Inches(5.5))
tf = content_box.text_frame
tf.word_wrap = True

# Estado
p = tf.paragraphs[0]
p.text = "Estado Final"
p.font.size = Pt(24)
p.font.bold = True
p.font.color.rgb = COLOR_SUCCESS
p.space_after = Pt(12)

p = tf.add_paragraph()
p.text = "✓ 5/5 xStocks completamente compatibles con la infraestructura API actual"
p.font.size = Pt(18)
p.level = 1
p.space_after = Pt(20)

# Conclusión
p = tf.add_paragraph()
p.text = "Conclusión"
p.font.size = Pt(24)
p.font.bold = True
p.font.color.rgb = COLOR_PRIMARY
p.space_after = Pt(12)

p = tf.add_paragraph()
p.text = "Todos los xStocks activos son 100% compatibles"
p.font.size = Pt(18)
p.level = 1

p = tf.add_paragraph()
p.text = "No se requieren cambios en el código"
p.font.size = Pt(18)
p.level = 1

p = tf.add_paragraph()
p.text = "La implementación puede proceder inmediatamente"
p.font.size = Pt(18)
p.level = 1

# ============================================================================
# SLIDE 3: xStocks Validados
# ============================================================================
slide = add_content_slide(prs, "xStocks Activos en Bybit")

# Tabla
rows = 6
cols = 4
left = Inches(1)
top = Inches(1.8)
width = Inches(8)
height = Inches(4.5)

table = slide.shapes.add_table(rows, cols, left, top, width, height).table

# Encabezados
headers = ['Símbolo', 'Nombre', 'Multiplicador', 'Estado']
for i, header in enumerate(headers):
    cell = table.cell(0, i)
    cell.text = header
    cell.fill.solid()
    cell.fill.fore_color.rgb = COLOR_PRIMARY
    paragraph = cell.text_frame.paragraphs[0]
    paragraph.font.bold = True
    paragraph.font.size = Pt(14)
    paragraph.font.color.rgb = RGBColor(255, 255, 255)
    paragraph.alignment = PP_ALIGN.CENTER

# Datos
data = [
    ['AAPLXUSDT', 'Apple Inc.', '1.0033', '✓ Trading'],
    ['NVDAXUSDT', 'NVIDIA Corp.', '1.0017', '✓ Trading'],
    ['METAXUSDT', 'Meta Platforms', '1.0029', '✓ Trading'],
    ['GOOGLXUSDT', 'Alphabet Inc.', '1.0024', '✓ Trading'],
    ['MCDXUSDT', "McDonald's", '1.0212', '✓ Trading']
]

for i, row_data in enumerate(data, start=1):
    for j, value in enumerate(row_data):
        cell = table.cell(i, j)
        cell.text = value
        paragraph = cell.text_frame.paragraphs[0]
        paragraph.font.size = Pt(12)
        paragraph.alignment = PP_ALIGN.CENTER

        # Color de fondo alternado
        if i % 2 == 0:
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(240, 240, 240)

        # Estado en verde
        if j == 3:
            paragraph.font.color.rgb = COLOR_SUCCESS
            paragraph.font.bold = True

# ============================================================================
# SLIDE 4: Problema Original
# ============================================================================
slide = add_content_slide(prs, "Problema Identificado")

content_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(9), Inches(5.5))
tf = content_box.text_frame
tf.word_wrap = True

# Problema
p = tf.paragraphs[0]
p.text = "Fallo en Validación Inicial"
p.font.size = Pt(24)
p.font.bold = True
p.font.color.rgb = COLOR_DANGER
p.space_after = Pt(12)

p = tf.add_paragraph()
p.text = "0/8 símbolos pasaron las pruebas de validación"
p.font.size = Pt(18)
p.level = 1
p.space_after = Pt(20)

# Causa raíz
p = tf.add_paragraph()
p.text = "Causa Raíz Identificada"
p.font.size = Pt(24)
p.font.bold = True
p.font.color.rgb = COLOR_PRIMARY
p.space_after = Pt(12)

p = tf.add_paragraph()
p.text = "Formato de símbolo incorrecto: se usaba 'AAPLX' en lugar de 'AAPLXUSDT'"
p.font.size = Pt(18)
p.level = 1

p = tf.add_paragraph()
p.text = "3 símbolos de la lista original no existen o no están activos"
p.font.size = Pt(18)
p.level = 1

p = tf.add_paragraph()
p.text = "Lógica de validación con errores en extracción de datos API"
p.font.size = Pt(18)
p.level = 1

# ============================================================================
# SLIDE 5: Proceso de Diagnóstico
# ============================================================================
slide = add_content_slide(prs, "Proceso de Diagnóstico")

content_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(9), Inches(5.5))
tf = content_box.text_frame
tf.word_wrap = True

steps = [
    ("1. Análisis de Resultados Fallidos", "Revisión del archivo xstocks_validation_results.json"),
    ("2. Script de Diagnóstico", "Creación de diagnose_xstocks.py con 11 variaciones de formato"),
    ("3. Descubrimiento de Formato Correcto", "Identificación de formato completo 'SYMBOLXUSDT'"),
    ("4. Listado Completo de Mercado", "Descarga de 527 símbolos spot de Bybit"),
    ("5. Filtrado de xStocks", "Identificación de 5 xStocks activos (no 8)"),
    ("6. Corrección de Validación", "Ajuste de lógica de extracción de datos API"),
    ("7. Validación Exitosa", "100% de compatibilidad confirmada")
]

for i, (title, desc) in enumerate(steps):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.text = title
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    p.space_after = Pt(6)

    p = tf.add_paragraph()
    p.text = desc
    p.font.size = Pt(14)
    p.level = 1
    p.space_after = Pt(12)

# ============================================================================
# SLIDE 6: Problemas Técnicos Resueltos
# ============================================================================
slide = add_content_slide(prs, "Problemas Técnicos Resueltos")

content_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(9), Inches(5.5))
tf = content_box.text_frame
tf.word_wrap = True

issues = [
    ("Error de Codificación Unicode",
     "Windows console no soporta caracteres ✓, ✗, ℹ",
     "Reemplazo por [OK], [ERROR], [INFO]"),

    ("Formato de Símbolo",
     "Se probaba 'AAPLX' en lugar de 'AAPLXUSDT'",
     "Corrección a formato completo del par"),

    ("Validación de BaseCoin",
     "Lógica incorrecta: esperaba 'AAPLXUSDT' en lugar de 'AAPLX'",
     "Extracción correcta eliminando 'USDT'"),

    ("Estructura JSON de Respuesta",
     "Category estaba en result, no en instrument",
     "Ajuste de extracción de datos anidados"),

    ("Mapeo de Valores",
     "Nombres de checks no coincidían con claves del objeto",
     "Creación de value_map para mapeo correcto")
]

for i, (problem, cause, solution) in enumerate(issues):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.text = f"{i+1}. {problem}"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = COLOR_DANGER
    p.space_after = Pt(4)

    p = tf.add_paragraph()
    p.text = f"Causa: {cause}"
    p.font.size = Pt(12)
    p.level = 1
    p.space_after = Pt(2)

    p = tf.add_paragraph()
    p.text = f"Solución: {solution}"
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_SUCCESS
    p.level = 1
    p.space_after = Pt(10)

# ============================================================================
# SLIDE 7: Pruebas de Compatibilidad API
# ============================================================================
slide = add_content_slide(prs, "Validación de Compatibilidad API")

content_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(9), Inches(5.5))
tf = content_box.text_frame
tf.word_wrap = True

p = tf.paragraphs[0]
p.text = "Todas las pruebas pasaron exitosamente para los 5 xStocks:"
p.font.size = Pt(18)
p.space_after = Pt(20)

tests = [
    ("✓ Instruments Info",
     ["Symbol correcto", "Category = spot", "BaseCoin y QuoteCoin correctos",
      "Campo xstockMultiplier presente"]),

    ("✓ Kline Data",
     ["Datos históricos disponibles", "Formato de velas correcto (7 elementos)",
      "Lógica OHLC válida"]),

    ("✓ Ticker Data",
     ["Precio en tiempo real", "Datos 24h completos", "Bid/Ask/Spread funcional"]),

    ("✓ Multiple Intervals",
     ["6 intervalos soportados: 1m, 5m, 15m, 60m, 240m, D",
      "Todos funcionan correctamente"])
]

for test_name, checks in tests:
    p = tf.add_paragraph()
    p.text = test_name
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_SUCCESS
    p.space_after = Pt(6)

    for check in checks:
        p = tf.add_paragraph()
        p.text = check
        p.font.size = Pt(14)
        p.level = 1
        p.space_after = Pt(2)

    tf.add_paragraph().space_after = Pt(10)

# ============================================================================
# SLIDE 8: Datos de Mercado Actual
# ============================================================================
slide = add_content_slide(prs, "Snapshot de Mercado (14 Oct 2026, 23:00 UTC)")

# Tabla
rows = 6
cols = 6
left = Inches(0.3)
top = Inches(1.8)
width = Inches(9.4)
height = Inches(4.5)

table = slide.shapes.add_table(rows, cols, left, top, width, height).table

# Encabezados
headers = ['Símbolo', 'Precio', '24h High', '24h Low', 'Vol 24h', 'Spread']
for i, header in enumerate(headers):
    cell = table.cell(0, i)
    cell.text = header
    cell.fill.solid()
    cell.fill.fore_color.rgb = COLOR_PRIMARY
    paragraph = cell.text_frame.paragraphs[0]
    paragraph.font.bold = True
    paragraph.font.size = Pt(13)
    paragraph.font.color.rgb = RGBColor(255, 255, 255)
    paragraph.alignment = PP_ALIGN.CENTER

# Datos
market_data = [
    ['AAPLXUSDT', '$347.16', '$348.51', '$342.98', '256', '0.029%'],
    ['NVDAXUSDT', '$220.21', '$230.38', '$218.38', '2,440', '0.032%'],
    ['METAXUSDT', '$730.64', '$770.82', '$729.02', '524', '0.036%'],
    ['GOOGLXUSDT', '$343.73', '$360.15', '$342.47', '414', '0.038%'],
    ['MCDXUSDT', '$240.58', '$244.00', '$236.00', '32', '0.224%']
]

for i, row_data in enumerate(market_data, start=1):
    for j, value in enumerate(row_data):
        cell = table.cell(i, j)
        cell.text = value
        paragraph = cell.text_frame.paragraphs[0]
        paragraph.font.size = Pt(11)
        paragraph.alignment = PP_ALIGN.CENTER

        # Color de fondo alternado
        if i % 2 == 0:
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(240, 240, 240)

# ============================================================================
# SLIDE 9: Archivos del Proyecto
# ============================================================================
slide = add_content_slide(prs, "Archivos Generados/Modificados")

content_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(9), Inches(5.5))
tf = content_box.text_frame
tf.word_wrap = True

files = [
    ("diagnose_xstocks.py", "Script de diagnóstico con 11 variaciones de formato"),
    ("filter_xstocks.py", "Filtro para extraer xStocks del listado completo"),
    ("debug_xstocks.py", "Script para ver respuesta RAW de la API"),
    ("xstocks_found.json", "527 símbolos spot con metadata (5 xStocks identificados)"),
    ("test_xstocks_api.py", "Script de validación corregido"),
    ("xstocks_validation_results.json", "Resultados finales exitosos")
]

p = tf.paragraphs[0]
p.text = "Scripts de Diagnóstico"
p.font.size = Pt(20)
p.font.bold = True
p.font.color.rgb = COLOR_PRIMARY
p.space_after = Pt(12)

for filename, description in files:
    p = tf.add_paragraph()
    p.text = filename
    p.font.size = Pt(14)
    p.font.bold = True
    p.space_after = Pt(4)

    p = tf.add_paragraph()
    p.text = description
    p.font.size = Pt(12)
    p.level = 1
    p.space_after = Pt(10)

# ============================================================================
# SLIDE 10: Estructura API Response
# ============================================================================
slide = add_content_slide(prs, "Estructura de Respuesta API (Ejemplo)")

content_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(9), Inches(5.5))
tf = content_box.text_frame
tf.word_wrap = True

code_example = '''GET /v5/market/instruments-info?category=spot&symbol=AAPLXUSDT

Response:
{
  "retCode": 0,
  "result": {
    "category": "spot",
    "list": [{
      "symbol": "AAPLXUSDT",
      "baseCoin": "AAPLX",
      "quoteCoin": "USDT",
      "status": "Trading",
      "symbolType": "xstocks",
      "xstockMultiplier": "1.0032690125398187",
      "lotSizeFilter": {...},
      "priceFilter": {...}
    }]
  }
}'''

p = tf.paragraphs[0]
p.text = code_example
p.font.name = 'Courier New'
p.font.size = Pt(11)
p.font.color.rgb = RGBColor(0, 0, 0)

# Agregar cuadro de fondo
shape = slide.shapes.add_shape(1, Inches(0.3), Inches(1.3), Inches(9.4), Inches(5.9))
shape.fill.solid()
shape.fill.fore_color.rgb = RGBColor(245, 245, 245)
shape.line.color.rgb = COLOR_PRIMARY
shape.line.width = Pt(2)

# Mover el text box al frente
slide.shapes._spTree.remove(content_box._element)
slide.shapes._spTree.append(content_box._element)

# ============================================================================
# SLIDE 11: Conclusiones y Próximos Pasos
# ============================================================================
slide = add_content_slide(prs, "Conclusiones y Próximos Pasos")

content_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(9), Inches(5.5))
tf = content_box.text_frame
tf.word_wrap = True

# Conclusiones
p = tf.paragraphs[0]
p.text = "Conclusiones Clave"
p.font.size = Pt(24)
p.font.bold = True
p.font.color.rgb = COLOR_SUCCESS
p.space_after = Pt(12)

conclusions = [
    "100% de compatibilidad confirmada para todos los xStocks activos",
    "API de Bybit funciona idénticamente para xStocks y pares spot tradicionales",
    "No se requieren cambios arquitectónicos en el código existente",
    "Formato correcto: símbolo completo 'SYMBOLXUSDT' (ej: AAPLXUSDT)"
]

for conclusion in conclusions:
    p = tf.add_paragraph()
    p.text = conclusion
    p.font.size = Pt(16)
    p.level = 1
    p.space_after = Pt(8)

# Próximos pasos
p = tf.add_paragraph()
p.text = "Próximos Pasos Recomendados"
p.font.size = Pt(24)
p.font.bold = True
p.font.color.rgb = COLOR_PRIMARY
p.space_after = Pt(12)

next_steps = [
    "Agregar los 5 símbolos xStocks a la configuración de símbolos",
    "Implementar en el entorno de desarrollo/staging",
    "Realizar pruebas de integración end-to-end",
    "Monitorear volumen y spreads antes del lanzamiento en producción",
    "Considerar expandir a más xStocks cuando Bybit los liste"
]

for step in next_steps:
    p = tf.add_paragraph()
    p.text = step
    p.font.size = Pt(16)
    p.level = 1
    p.space_after = Pt(8)

# ============================================================================
# SLIDE 12: Slide Final
# ============================================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank

# Fondo
background = slide.background
fill = background.fill
fill.solid()
fill.fore_color.rgb = COLOR_PRIMARY

# Mensaje
message_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(9), Inches(1.5))
message_frame = message_box.text_frame
message_frame.text = "Validación Completada"
message_p = message_frame.paragraphs[0]
message_p.alignment = PP_ALIGN.CENTER
message_p.font.size = Pt(54)
message_p.font.bold = True
message_p.font.color.rgb = RGBColor(255, 255, 255)

# Submensaje
sub_box = slide.shapes.add_textbox(Inches(0.5), Inches(4.2), Inches(9), Inches(1))
sub_frame = sub_box.text_frame
sub_frame.text = "xStocks Ready for Implementation"
sub_p = sub_frame.paragraphs[0]
sub_p.alignment = PP_ALIGN.CENTER
sub_p.font.size = Pt(28)
sub_p.font.color.rgb = COLOR_SUCCESS

# Fecha
date_box = slide.shapes.add_textbox(Inches(0.5), Inches(6.5), Inches(9), Inches(0.5))
date_frame = date_box.text_frame
date_frame.text = "Octubre 2026"
date_p = date_frame.paragraphs[0]
date_p.alignment = PP_ALIGN.CENTER
date_p.font.size = Pt(18)
date_p.font.color.rgb = RGBColor(200, 200, 200)

# ============================================================================
# Guardar presentación
# ============================================================================
prs.save('xStocks_Validation_Report.pptx')
print("[OK] Presentacion creada exitosamente: xStocks_Validation_Report.pptx")
print(f"[INFO] Total de slides: {len(prs.slides)}")
